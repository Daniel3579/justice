import io
import re
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from jinja2 import Environment, FileSystemLoader, StrictUndefined

from calc import calc_penalty, fmt_date, money, refund_amount
from config import ISK_TEMPLATE, TEMPLATES_PATH
from schemas import CaseFacts, Composed

NEED_RE = re.compile(r"\[ТРЕБУЕТСЯ УТОЧНЕНИЕ:\s*(.+?)\]")
TAG_RE = re.compile(r"\s*\[(?:НОРМА|WEB)-\d+\]")
MARKER_RE = re.compile(r"^\[(R|C|B|L)\]")
DUTY_LIMIT = 1_000_000

# Здесь формируется сам документы 
def need(value, label: str):
    if value is None or (isinstance(value, str) and not value.strip()):
        return f"[ТРЕБУЕТСЯ УТОЧНЕНИЕ: {label}]"
    return value


_env = Environment(
    loader=FileSystemLoader(str(TEMPLATES_PATH)),
    undefined=StrictUndefined,
    autoescape=False,
    trim_blocks=True,
    lstrip_blocks=True,
    keep_trailing_newline=False,
)
_env.filters["need"] = need


def short_name(full: str | None) -> str | None:
    if not full:
        return None
    parts = full.split()
    if len(parts) < 2:
        return full
    initials = "".join(f"{p[0]}." for p in parts[1:3])
    return f"{parts[0]} {initials}"


def strip_tags(text: str) -> str:
    return TAG_RE.sub("", text).strip()


def clean_paragraphs(items: list[str], noise: list[str]) -> list[str]:
    result = []
    for item in items:
        for line in item.splitlines():
            line = strip_tags(line)
            for name in noise:
                line = line.replace(name, "")
            line = re.sub(r"\s{2,}", " ", line).strip(" ,;:-—")
            if len(line) < 8 or line.lower().startswith("источник") or ".pdf" in line.lower():
                continue
            result.append(line)
    return result


def describe_contract(facts: CaseFacts) -> str | None:
    parts = []
    if facts.contract_number:
        parts.append(f"№ {facts.contract_number}")
    if facts.contract_date:
        parts.append(f"от {fmt_date(facts.contract_date)}")
    return " ".join(parts) or None


def sentence(text: str) -> str:
    text = text.strip()
    text = text[:1].upper() + text[1:]
    return text if text.endswith(".") else text + "."


def build_demands(facts: CaseFacts, refund: float | None, penalty: dict, contract: str | None) -> list[str]:
    demands = []
    by_contract = " по договору" + (f" {contract}" if contract else "")

    if refund:
        demands.append(f"Взыскать с ответчика в пользу истца денежные средства, уплаченные{by_contract}, в размере {money(refund)}")
    elif facts.violation_type != "other" and not facts.extra_demands:
        demands.append("Взыскать с ответчика в пользу истца денежные средства, уплаченные по договору, "
                       "в размере [ТРЕБУЕТСЯ УТОЧНЕНИЕ: сумма, уплаченная по договору].")

    if penalty.get("available"):
        demands.append(f"Взыскать с ответчика в пользу истца неустойку за период с {penalty['start']} "
                       f"по {penalty['end']} в размере {penalty['amount_str']}")
    elif penalty.get("applicable"):
        reasons = "; ".join(penalty["missing"])
        demands.append(f"Взыскать с ответчика в пользу истца неустойку на основании {penalty['basis']} "
                       f"[ТРЕБУЕТСЯ УТОЧНЕНИЕ: {reasons}]")

    demands.extend(facts.extra_demands)

    moral = money(facts.moral_damage) if facts.moral_damage else "[ТРЕБУЕТСЯ УТОЧНЕНИЕ: сумма компенсации морального вреда]"
    demands.append(f"Взыскать с ответчика в пользу истца компенсацию морального вреда в размере {moral}")
    demands.append("Взыскать с ответчика в пользу истца штраф за несоблюдение в добровольном порядке "
                   "удовлетворения требований потребителя в размере 50% от суммы, присуждённой судом в пользу истца")
    return [sentence(d) for d in demands]


def build_attachments(facts: CaseFacts) -> list[str]:
    items = ["Документ, подтверждающий направление ответчику копии искового заявления и приложенных к нему документов"]
    if facts.contract_type or describe_contract(facts):
        items.append("Копия договора")
    if facts.paid_amount or facts.price:
        items.append("Копии документов, подтверждающих оплату")
    if facts.claim_sent or facts.claim_date:
        items.append("Копия претензии и доказательства её направления ответчику")
    if facts.claim_response:
        items.append("Копия ответа ответчика на претензию")
    items.extend(facts.documents)
    items.append("Копия документа, удостоверяющего личность истца")

    seen, result = set(), []
    for item in items:
        item = sentence(item)
        key = item.lower()
        if key not in seen:
            seen.add(key)
            result.append(item)
    return result


def build_context(facts: CaseFacts, composed: Composed, noise: list[str] | None = None, today: date | None = None) -> dict:
    today = today or date.today()
    penalty = calc_penalty(facts, today)
    refund = refund_amount(facts)
    contract = describe_contract(facts)

    total = (refund or 0) + (penalty.get("amount") or 0)
    claim_price = money(total) if total else "[ТРЕБУЕТСЯ УТОЧНЕНИЕ: цена иска]"
    duty_exempt = total <= DUTY_LIMIT

    plaintiff = facts.plaintiff.model_dump()
    plaintiff["short"] = short_name(facts.plaintiff.name)

    return {
        "court": facts.court,
        "plaintiff": plaintiff,
        "defendant": facts.defendant.model_dump(),
        "claim_price": claim_price,
        "duty_exempt": duty_exempt,
        "facts_text": clean_paragraphs(composed.facts_text, noise or []) or ["[ТРЕБУЕТСЯ УТОЧНЕНИЕ: обстоятельства дела]"],
        "legal_text": clean_paragraphs(composed.legal_text, noise or []),
        "penalty": penalty,
        "moral_damage": money(facts.moral_damage) if facts.moral_damage else None,
        "demands": build_demands(facts, refund, penalty, contract),
        "attachments": build_attachments(facts),
        "date": fmt_date(today),
    }


def render_marked(context: dict) -> str:
    text = _env.get_template(ISK_TEMPLATE).render(**context)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def strip_markers(marked: str) -> str:
    return "\n".join(MARKER_RE.sub("", line) for line in marked.splitlines())


def find_missing(text: str) -> list[str]:
    seen, result = set(), []
    for item in NEED_RE.findall(text):
        if item not in seen:
            seen.add(item)
            result.append(item)
    return result


def _set_font(run, bold: bool = False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")


def marked_to_docx(marked: str) -> bytes:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")

    for line in marked.splitlines():
        m = MARKER_RE.match(line)
        marker = m.group(1) if m else None
        text = MARKER_RE.sub("", line).rstrip()

        p = doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.15

        if marker == "R":
            fmt.left_indent = Cm(8)
            fmt.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            bold = False
        elif marker == "C":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            bold = True
        elif marker == "B":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            bold = True
        elif marker == "L":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            bold = False
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            fmt.first_line_indent = Cm(1.25) if text else None
            bold = False

        if text:
            _set_font(p.add_run(text), bold)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()